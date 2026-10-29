from __future__ import annotations

import csv
import io
import json
import smtplib
from email.message import EmailMessage
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

from convoke.audit import AuditLog, utcnow
from convoke.config import Settings
from convoke.paths import NOTICES, OUTPUTS, ensure_dirs
from convoke.security import sanitize_untrusted
from convoke.store import execute, fetchall, fetchone, log_action, sha256_bytes

audit = AuditLog()

STATUSES = {
    "Not started": "Registered, no current session seat.",
    "In progress": "Waitlisted or pending a missing preference.",
    "Evidence supports completion": "Confirmed seat with capacity remaining after checks.",
    "Delayed": "Cancellation or timetable change still being repaired.",
    "Needs review": "Conflict, missing evidence, or over-capacity claim.",
}


def _csv_rows(text: str) -> list[dict[str, str]]:
    handle = io.StringIO(text)
    return list(csv.DictReader(handle))


def read_source(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(raw))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages), sha256_bytes(raw)
    text = raw.decode("utf-8", errors="replace")
    return text, sha256_bytes(raw)


def parse_registration_rows(path: str) -> list[dict[str, str]]:
    source = Path(path)
    text, _digest = read_source(source)
    sanitised = sanitize_untrusted(text)
    if source.suffix.lower() == ".json":
        payload = json.loads(sanitised.cleaned or "[]")
        return payload if isinstance(payload, list) else payload.get("registrations", [])
    try:
        return _csv_rows(sanitised.cleaned)
    except Exception:
        return []


def tool_policy_search(query: str) -> dict[str, Any]:
    """Retrieve approved local policy passages. Source text is untrusted except this file."""
    from convoke.paths import SAMPLES

    path = SAMPLES / "policy.txt"
    text = path.read_text(encoding="utf-8")
    needles = [part.strip().lower() for part in query.split() if len(part.strip()) > 3]
    lines = text.splitlines()
    hits = []
    for index, line in enumerate(lines, start=1):
        low = line.lower()
        if any(needle in low for needle in needles) or line.startswith("-"):
            hits.append({"line": index, "text": line})
    log_action(None, None, "tool", "policy_search", {"query": query, "hits": len(hits)}, "ok")
    return {"ok": True, "source": str(path), "passages": hits[:8]}


def tool_registration_reader(event_id: str, cycle_id: int, path: str) -> dict[str, Any]:
    """Ingest a registration file. Duplicate hashes are skipped."""
    source = Path(path)
    text, digest = read_source(source)
    existing = fetchone("SELECT * FROM processed_sources WHERE source_hash=?", (digest,))
    if existing:
        log_action(event_id, cycle_id, "tool", "registration_reader", {"path": path, "hash": digest}, "duplicate_skipped")
        audit.write("registration_reader", outcome="duplicate_skipped", hash=digest, path=path)
        return {"ok": True, "duplicate": True, "hash": digest, "inserted": 0, "flagged": existing["flagged_injection"]}

    sanitised = sanitize_untrusted(text)
    rows = []
    if source.suffix.lower() == ".json":
        payload = json.loads(sanitised.cleaned)
        rows = payload if isinstance(payload, list) else payload.get("registrations", [])
    else:
        # CSV, or PDF/text that still contains a CSV table
        try:
            rows = _csv_rows(sanitised.cleaned)
        except Exception:
            rows = []

    inserted = 0
    for row in rows:
        pid = (row.get("participant_id") or row.get("id") or "").strip()
        name = (row.get("name") or row.get("display_name") or "").strip()
        if not pid or not name:
            continue
        execute(
            """
            INSERT INTO participants (participant_id, event_id, display_name, email, preference, dietary, source_file, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(event_id, participant_id) DO UPDATE SET
                display_name=excluded.display_name,
                email=excluded.email,
                preference=excluded.preference,
                dietary=excluded.dietary,
                source_file=excluded.source_file
            """,
            (
                pid,
                event_id,
                name,
                (row.get("email") or "").strip(),
                (row.get("preference") or row.get("session") or "").strip(),
                (row.get("dietary") or "").strip(),
                source.name,
                utcnow(),
            ),
        )
        inserted += 1

    execute(
        """
        INSERT INTO processed_sources (source_hash, event_id, filename, kind, flagged_injection, processed_at)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (digest, event_id, source.name, source.suffix, int(sanitised.flagged), utcnow()),
    )
    log_action(
        event_id,
        cycle_id,
        "tool",
        "registration_reader",
        {"path": path, "hash": digest, "flagged": sanitised.flagged, "matches": sanitised.matches},
        f"inserted:{inserted}",
    )
    audit.write("registration_reader", hash=digest, inserted=inserted, flagged=sanitised.flagged)
    return {
        "ok": True,
        "duplicate": False,
        "hash": digest,
        "inserted": inserted,
        "flagged": sanitised.flagged,
        "matches": sanitised.matches,
    }


def tool_memory_retrieve(event_id: str, cycle_id: int, enabled: bool = True) -> dict[str, Any]:
    if not enabled:
        log_action(event_id, cycle_id, "memory", "memory_retrieve", {"enabled": False}, "disabled")
        return {"enabled": False, "sessions": [], "participants": [], "allocations": [], "notices": [], "working": None}

    sessions = fetchall("SELECT * FROM sessions WHERE event_id=? AND is_current=1", (event_id,))
    participants = fetchall("SELECT * FROM participants WHERE event_id=?", (event_id,))
    allocations = fetchall(
        "SELECT * FROM allocations WHERE event_id=? AND is_current=1 ORDER BY cycle_id DESC, allocation_id DESC",
        (event_id,),
    )
    waitlist = fetchall("SELECT * FROM waitlist WHERE event_id=? AND is_current=1 ORDER BY position", (event_id,))
    notices = fetchall("SELECT * FROM notices WHERE event_id=? ORDER BY notice_id DESC LIMIT 50", (event_id,))
    working = fetchone("SELECT * FROM working_state WHERE event_id=?", (event_id,))
    history = fetchall(
        "SELECT * FROM session_history WHERE event_id=? ORDER BY changed_at DESC LIMIT 20",
        (event_id,),
    )
    log_action(event_id, cycle_id, "memory", "memory_retrieve", {"counts": {
        "sessions": len(sessions),
        "participants": len(participants),
        "allocations": len(allocations),
    }}, "ok")
    return {
        "enabled": True,
        "sessions": sessions,
        "participants": participants,
        "allocations": allocations,
        "waitlist": waitlist,
        "notices": notices,
        "working": working,
        "history": history,
    }


def tool_capacity_calculator(event_id: str, cycle_id: int, requests: list[dict[str, str]]) -> dict[str, Any]:
    """Deterministic seat math. The LLM never counts remaining seats."""
    sessions = {
        row["session_id"]: dict(row)
        for row in fetchall("SELECT * FROM sessions WHERE event_id=? AND is_current=1", (event_id,))
    }
    counts = {sid: 0 for sid in sessions}
    result = []
    waitlist_plan = []

    for req in requests:
        pid = req["participant_id"]
        wanted = req.get("requested_session") or ""
        action = req.get("action", "place")
        if action == "cancel":
            result.append(
                {
                    "participant_id": pid,
                    "session_id": None,
                    "agent_status": "Delayed",
                    "blocker": "Cancellation recorded; seat released.",
                    "next_action": "promote_waitlist",
                    "evidence": "cancellation request",
                }
            )
            continue
        if wanted not in sessions:
            result.append(
                {
                    "participant_id": pid,
                    "session_id": None,
                    "agent_status": "Needs review",
                    "blocker": "Requested session is missing or ambiguous.",
                    "next_action": "ask_clarification",
                    "evidence": f"request={wanted!r}",
                }
            )
            continue
        cap = int(sessions[wanted]["capacity"])
        if counts[wanted] < cap:
            counts[wanted] += 1
            result.append(
                {
                    "participant_id": pid,
                    "session_id": wanted,
                    "agent_status": "Evidence supports completion",
                    "blocker": None,
                    "next_action": "draft_confirmation",
                    "evidence": f"{wanted} seats used {counts[wanted]}/{cap}",
                }
            )
        else:
            waitlist_plan.append({"session_id": wanted, "participant_id": pid})
            result.append(
                {
                    "participant_id": pid,
                    "session_id": wanted,
                    "agent_status": "In progress",
                    "blocker": "Session at capacity.",
                    "next_action": "waitlist_and_notify",
                    "evidence": f"{wanted} full at {cap}",
                }
            )

    remaining = {sid: int(sessions[sid]["capacity"]) - counts[sid] for sid in sessions}
    placed_ids = {item["participant_id"] for item in result}
    previous_wait = fetchall(
        "SELECT * FROM waitlist WHERE event_id=? AND is_current=1 ORDER BY position",
        (event_id,),
    )
    promoted = []
    for item in previous_wait:
        pid = item["participant_id"]
        sid = item["session_id"]
        if pid in placed_ids or sid not in sessions:
            continue
        already = next((row for row in result if row["participant_id"] == pid), None)
        if already and already.get("agent_status") == "Delayed":
            continue
        if remaining.get(sid, 0) <= 0:
            continue
        counts[sid] += 1
        remaining[sid] = int(sessions[sid]["capacity"]) - counts[sid]
        waitlist_plan = [w for w in waitlist_plan if w["participant_id"] != pid]
        promoted.append(pid)
        if already:
            already.update(
                {
                    "session_id": sid,
                    "agent_status": "Evidence supports completion",
                    "blocker": None,
                    "next_action": "draft_confirmation",
                    "evidence": f"promoted from waitlist into {sid}; seats {counts[sid]}/{sessions[sid]['capacity']}",
                }
            )
        else:
            result.append(
                {
                    "participant_id": pid,
                    "session_id": sid,
                    "agent_status": "Evidence supports completion",
                    "blocker": None,
                    "next_action": "draft_confirmation",
                    "evidence": f"promoted from waitlist into {sid}",
                }
            )
            placed_ids.add(pid)
    remaining = {sid: int(sessions[sid]["capacity"]) - counts[sid] for sid in sessions}
    log_action(
        event_id,
        cycle_id,
        "tool",
        "capacity_calculator",
        {"remaining": remaining, "promoted": promoted},
        "ok",
    )
    return {
        "ok": True,
        "placements": result,
        "waitlist": waitlist_plan,
        "remaining": remaining,
        "promoted": promoted,
        "status_defs": STATUSES,
    }


def tool_record_allocations(
    event_id: str,
    cycle_id: int,
    placements: list[dict[str, Any]],
    waitlist: list[dict[str, str]],
    simulate_failure: bool = False,
) -> dict[str, Any]:
    if simulate_failure:
        log_action(event_id, cycle_id, "tool", "record_allocations", {"simulate_failure": True}, "failed", "simulated write failure")
        return {"ok": False, "error": "simulated write failure", "written": 0}

    execute(
        "UPDATE allocations SET is_current=0 WHERE event_id=? AND is_current=1 AND cycle_id<?",
        (event_id, cycle_id),
    )
    written = 0
    for item in placements:
        existing = fetchone(
            "SELECT allocation_id FROM allocations WHERE event_id=? AND cycle_id=? AND participant_id=?",
            (event_id, cycle_id, item["participant_id"]),
        )
        if existing:
            continue
        participant = fetchone(
            "SELECT preference FROM participants WHERE participant_id=? AND event_id=?",
            (item["participant_id"], event_id),
        )
        execute(
            """
            INSERT INTO allocations (
                event_id, cycle_id, session_id, participant_id, requested_session,
                agent_status, organizer_status, organizer_comment, evidence, blocker,
                next_action, last_contact, is_current, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?)
            """,
            (
                event_id,
                cycle_id,
                item.get("session_id"),
                item["participant_id"],
                (participant or {}).get("preference"),
                item["agent_status"],
                None,
                None,
                item.get("evidence"),
                item.get("blocker"),
                item.get("next_action"),
                "none",
                utcnow(),
            ),
        )
        written += 1

    execute("UPDATE waitlist SET is_current=0 WHERE event_id=?", (event_id,))
    for index, item in enumerate(waitlist, start=1):
        execute(
            """
            INSERT INTO waitlist (event_id, session_id, participant_id, position, reason, is_current, created_at)
            VALUES (?, ?, ?, ?, ?, 1, ?)
            """,
            (event_id, item["session_id"], item["participant_id"], index, "capacity", utcnow()),
        )
    log_action(event_id, cycle_id, "memory", "record_allocations", {"written": written}, "ok")
    return {"ok": True, "written": written, "waitlisted": len(waitlist)}


def tool_update_session_fact(
    event_id: str,
    session_id: str,
    field: str,
    new_value: str,
    approved_by: str,
    reason: str,
) -> dict[str, Any]:
    current = fetchone(
        "SELECT * FROM sessions WHERE session_id=? AND event_id=? AND is_current=1",
        (session_id, event_id),
    )
    if not current:
        return {"ok": False, "error": "session not found"}
    allowed = {"capacity", "starts_at", "ends_at", "room", "name"}
    if field not in allowed:
        return {"ok": False, "error": f"field {field} cannot be updated"}
    old = str(current[field])
    if field == "capacity":
        execute(
            "UPDATE sessions SET capacity=?, version=version+1 WHERE session_id=? AND event_id=?",
            (int(new_value), session_id, event_id),
        )
    elif field == "starts_at":
        execute("UPDATE sessions SET starts_at=? WHERE session_id=? AND event_id=?", (new_value, session_id, event_id))
    elif field == "ends_at":
        execute("UPDATE sessions SET ends_at=? WHERE session_id=? AND event_id=?", (new_value, session_id, event_id))
    elif field == "room":
        execute("UPDATE sessions SET room=? WHERE session_id=? AND event_id=?", (new_value, session_id, event_id))
    else:
        execute("UPDATE sessions SET name=? WHERE session_id=? AND event_id=?", (new_value, session_id, event_id))
    execute(
        """
        INSERT INTO session_history (session_id, event_id, field, old_value, new_value, reason, approved_by, changed_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (session_id, event_id, field, old, str(new_value), reason, approved_by, utcnow()),
    )
    log_action(event_id, None, "memory", "update_session_fact", {"field": field, "old": old, "new": new_value}, "superseded")
    return {"ok": True, "old": old, "new": str(new_value)}


def _notice_key(event_id: str, cycle_id: int, participant_id: str, kind: str) -> str:
    from convoke.store import sha256_text

    return sha256_text(f"{event_id}|{cycle_id}|{participant_id}|{kind}")


def tool_notice_writer(
    event_id: str,
    cycle_id: int,
    participant_id: str,
    kind: str,
    subject: str,
    body: str,
    recipient: str,
    settings: Settings,
    send: bool = False,
    approved: bool = False,
    simulate_failure: bool = False,
) -> dict[str, Any]:
    ensure_dirs()
    key = _notice_key(event_id, cycle_id, participant_id, kind)
    existing = fetchone("SELECT * FROM notices WHERE idempotency_key=?", (key,))
    if existing:
        log_action(event_id, cycle_id, "tool", "notice_writer", {"key": key}, "duplicate_prevented")
        return {"ok": True, "duplicate": True, "status": existing["status"], "notice_id": existing["notice_id"]}

    status = "UNSENT"
    outcome = "draft_saved"
    if send and approved:
        allow = settings.allowlist
        dest = (recipient or "").lower()
        if allow and dest not in allow:
            status = "UNSENT"
            outcome = "blocked_not_allowlisted"
        elif simulate_failure:
            status = "FAILED"
            outcome = "simulated smtp failure"
        elif not settings.smtp_username or not settings.smtp_password:
            status = "UNSENT"
            outcome = "no_smtp_credentials"
        else:
            try:
                message = EmailMessage()
                message["Subject"] = subject
                message["From"] = settings.smtp_from or settings.smtp_username
                message["To"] = recipient
                message.set_content(body)
                with smtplib.SMTP(settings.smtp_host, settings.smtp_port, timeout=20) as smtp:
                    smtp.starttls()
                    smtp.login(settings.smtp_username, settings.smtp_password)
                    smtp.send_message(message)
                status = "SENT"
                outcome = "delivered"
            except Exception as exc:
                status = "FAILED"
                outcome = str(exc)

    notice_id = execute(
        """
        INSERT INTO notices (
            idempotency_key, event_id, participant_id, cycle_id, subject, body,
            recipient, status, delivery_outcome, approved, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (key, event_id, participant_id, cycle_id, subject, body, recipient, status, outcome, int(approved), utcnow()),
    )
    draft_path = NOTICES / f"{key[:12]}_{status}.txt"
    draft_path.write_text(f"TO: {recipient}\nSUBJECT: {subject}\nSTATUS: {status}\n\n{body}\n", encoding="utf-8")
    log_action(event_id, cycle_id, "tool", "notice_writer", {"key": key, "status": status}, outcome)
    return {"ok": True, "duplicate": False, "status": status, "outcome": outcome, "notice_id": notice_id, "path": str(draft_path)}


def tool_excel_register(event_id: str, preserve_manual: bool = True) -> str:
    ensure_dirs()
    path = OUTPUTS / f"{event_id}_register.xlsx"
    previous: dict[tuple, dict] = {}
    if preserve_manual and path.exists():
        book = load_workbook(path)
        sheet = book.active
        headers = [cell.value for cell in sheet[1]]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            record = dict(zip(headers, row))
            key = (record.get("event_id"), record.get("cycle_id"), record.get("participant_id"))
            previous[key] = record

    rows = fetchall(
        """
        SELECT a.*, p.display_name, p.email, s.name as session_name, s.capacity, s.starts_at
        FROM allocations a
        JOIN participants p ON p.participant_id = a.participant_id AND p.event_id = a.event_id
        LEFT JOIN sessions s ON s.session_id = a.session_id AND s.event_id = a.event_id AND s.is_current = 1
        WHERE a.event_id=?
        ORDER BY a.cycle_id, a.allocation_id
        """,
        (event_id,),
    )
    book = Workbook()
    sheet = book.active
    sheet.title = "progress_register"
    headers = [
        "event_id",
        "review_cycle_id",
        "session_id",
        "session_name",
        "approved_capacity",
        "session_start",
        "participant_id",
        "participant_name",
        "student_claim",
        "agent_suggested_status",
        "evidence_references",
        "missing_evidence_or_blocker",
        "next_action",
        "last_contact_outcome",
        "organizer_decision",
        "organizer_comments",
        "review_timestamp",
    ]
    sheet.append(headers)
    for row in rows:
        key = (row["event_id"], row["cycle_id"], row["participant_id"])
        prior = previous.get(key, {})
        sheet.append(
            [
                row["event_id"],
                row["cycle_id"],
                row["session_id"],
                row["session_name"],
                row["capacity"],
                row["starts_at"],
                row["participant_id"],
                row["display_name"],
                row["requested_session"],
                row["agent_status"],
                row["evidence"],
                row["blocker"],
                row["next_action"],
                row["last_contact"],
                prior.get("organizer_decision") or row["organizer_status"],
                prior.get("organizer_comments") or row["organizer_comment"],
                row["created_at"],
            ]
        )
    book.save(path)
    log_action(event_id, None, "tool", "excel_register", {"path": str(path)}, "ok")
    return str(path)
